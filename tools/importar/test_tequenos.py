from datetime import date, datetime, time

import openpyxl
import pytest

from extractores.tequenos import (ARCHIVOS, extraer, fechas_produccion, fusionar, hay_freido, lote_y_codigo, opcion,
                                  presentacion, producto, proveedor, quitar_repetidos, segundos, temp_agua,
                                  vencimiento)

HOY = date(2026, 10, 12)


# ───────────── lote ─────────────
def test_lote_con_codigo_entre_corchetes():
    assert lote_y_codigo('[PT011] 25131') == ('25131', 'PT011')


def test_lote_con_numero_entre_parentesis():
    assert lote_y_codigo('PT014 (27037)') == ('27037', 'PT014')
    assert lote_y_codigo('PT035(27131)') == ('27131', 'PT035')


def test_lote_conserva_el_codigo_tal_como_se_escribio():
    assert lote_y_codigo('PY026(28057)') == ('28057', 'PY026')


def test_lote_sin_dato():
    assert lote_y_codigo('-') == (None, None)
    assert lote_y_codigo(None) == (None, None)


def test_lote_numerico():
    assert lote_y_codigo(28861) == ('28861', None)


# ───────────── tiempo de freído ─────────────
def test_segundos_desde_hora_de_excel():
    assert segundos(time(0, 1, 8)) == 68
    assert segundos(time(0, 0, 57)) == 57
    assert segundos(time(0, 1)) == 60


def test_segundos_desde_texto():
    assert segundos('1 Minuto') == 60
    assert segundos('1 minuto') == 60
    assert segundos('90 s') == 90


def test_segundos_sin_dato():
    assert segundos('NA') is None
    assert segundos('-') is None
    assert segundos('rápido') is None


# ───────────── agua ─────────────
def test_temp_agua():
    assert temp_agua('Ambiente') == 'Ambiente'
    assert temp_agua(' ambiente ') == 'Ambiente'
    assert temp_agua(0.7) == 0.7
    assert temp_agua('-') is None


# ───────────── opciones ─────────────
def test_opcion_ignora_mayusculas_acentos_y_espacios():
    assert opcion('Ausencia ', ['Ausencia', 'Presencia']) == 'Ausencia'
    assert opcion('buena', ['Excelente', 'Buena', 'Regular', 'Mala']) == 'Buena'


def test_correcto_y_adecuada_quedan_como_buena_si_no_hay_otra_evaluacion():
    # Operaciones (12/10/2026): "Correcto"/"Adecuada" = Buena; lo descriptivo sigue en observaciones
    from extractores.tequenos import equivaler_escala
    x = {'valores': {'tMasa': 33.0}, 'pendientes': {'aspectoFreido': 'Correcto', 'sabor': 'adecuada', 'color': 'Amarillento'}}
    equivaler_escala(x)
    assert x['valores'] == {'tMasa': 33.0, 'aspectoFreido': 'Buena', 'sabor': 'Buena'}
    assert x['pendientes'] == {'color': 'Amarillento'}


def test_la_evaluacion_de_la_hoja_vieja_gana_sobre_correcto():
    from extractores.tequenos import equivaler_escala, fusionar
    principal = {'valores': {'aspectoFreido': None}, 'pendientes': {'aspectoFreido': 'Correcto'}}
    vieja = {'valores': {'aspectoFreido': 'Excelente'}, 'pendientes': {}}
    x = fusionar(principal, vieja)
    equivaler_escala(x)
    assert x['valores']['aspectoFreido'] == 'Excelente'
    assert x['pendientes'] == {}


def test_opcion_fuera_de_la_escala_es_none():
    assert opcion('Correcto', ['Excelente', 'Buena', 'Regular', 'Mala']) is None
    assert opcion('Amarillento', ['Excelente', 'Buena', 'Regular', 'Mala']) is None


# ───────────── textos ─────────────
def test_presentacion_uniforma_unidades():
    assert presentacion('24 UNIDADES') == '24 unidades'
    assert presentacion('12 UNIDADES ') == '12 unidades'
    assert presentacion('-') is None


def test_presentacion_deja_otros_textos():
    assert presentacion('Caja blanca') == 'Caja blanca'


def test_proveedor_capitaliza_palabras_en_minuscula():
    assert proveedor('quesos rafa') == 'Quesos Rafa'
    assert proveedor('quesos San Miguel') == 'Quesos San Miguel'
    assert proveedor('Quesos SM') == 'Quesos SM'
    assert proveedor('-') is None


# ───────────── fechas ─────────────
def test_vencimiento_elige_la_lectura_a_seis_meses_de_la_produccion():
    # Excel leyó 11/06/2026 como 6 de noviembre
    assert vencimiento(datetime(2026, 11, 6), date(2025, 12, 11)) == date(2026, 6, 11)
    assert vencimiento(datetime(2027, 4, 2), date(2026, 8, 4)) == date(2027, 2, 4)


def test_vencimiento_en_texto_es_exacto():
    assert vencimiento('29/06/2026', date(2026, 6, 29)) == date(2026, 6, 29)


def test_vencimiento_sin_produccion_o_sin_dato():
    assert vencimiento(datetime(2026, 11, 6), None) == date(2026, 11, 6)
    assert vencimiento('-', date(2026, 1, 1)) is None


def test_fechas_produccion_corrige_y_separa_las_fuera_de_rango():
    crudas = [datetime(2025, 11, 12), '13/12/2025', '15/04/2028', datetime(2026, 11, 4), None]
    fechas, fuera = fechas_produccion(crudas, HOY)
    assert fechas == [date(2025, 12, 11), date(2025, 12, 13), None, date(2026, 4, 11), None]
    assert fuera == {2: date(2028, 4, 15)}


def test_fechas_produccion_no_usa_fechas_futuras_como_referencia():
    # sin el filtro, la fecha ambigua se acercaría a 2028 y quedaría en el 4 de mayo
    crudas = ['15/04/2028', datetime(2026, 4, 5)]
    fechas, fuera = fechas_produccion(crudas, HOY)
    assert fechas[1] == date(2026, 4, 5)


# ───────────── producto ─────────────
def test_producto_full_queso_crudo_y_precocido():
    assert producto('FULL QUESO', 'PT011', '12 unidades', freido=False) == ('Tequeño Full Queso', None)
    assert producto('FULL QUESO', 'PT026', '24 unidades', freido=True) == ('Tequeño Full Queso precocido', None)


def test_producto_precocido_por_nombre_aunque_no_haya_freido():
    assert producto('Tequeño full queso prcocido', 'PT013', '20 unidades', freido=False)[0] == \
        'Tequeño Full Queso precocido'


def test_producto_full_de_20_y_24_es_precocido_aunque_no_haya_freido():
    # Operaciones (12/10/2026): las presentaciones de 20 y 24 unidades son precocidas
    assert producto('FULL QUESO', 'PT026', '24 unidades', freido=False)[0] == 'Tequeño Full Queso precocido'
    assert producto('FULL QUESO', 'PT013', '20 unidades', freido=False)[0] == 'Tequeño Full Queso precocido'
    assert producto('FULL QUESO', 'PT011', '12 unidades', freido=False)[0] == 'Tequeño Full Queso'


def test_producto_otros_de_la_linea():
    assert producto('BOQUITAS', 'PT006', '60 unidades', freido=False) == ('Tequeño Boquita', None)
    assert producto('Tequeñón', 'PT015', '50 unidades', freido=True) == ('Tequeñón', None)
    assert producto('Guayaba queso', 'PT029', '10 unidades', freido=True) == ('Tequeño Full de Guayaba (5 cm)', None)


def test_producto_salmon_por_presentacion():
    assert producto('Salmón', 'PT035', '10 unidades', freido=True) == ('PT035', None)
    assert producto('Salmón', None, '50 unidades', freido=False) == ('PT036', None)
    clave, motivo = producto('Salmón', None, None, freido=False)
    assert clave is None and motivo


def test_producto_nombre_y_codigo_de_familias_distintas_queda_sin_producto():
    clave, motivo = producto('BOQUITAS', 'PT026', '24 unidades', freido=False)
    assert clave is None
    assert 'PT026' in motivo


def test_producto_codigo_desconocido_no_contradice():
    assert producto('FULL QUESO', 'PY026', '24 unidades', freido=True)[0] == 'Tequeño Full Queso precocido'


def test_producto_nombre_desconocido():
    clave, motivo = producto('Mixto especial', None, None, freido=False)
    assert clave is None and motivo


def test_las_claves_que_produce_existen_en_el_catalogo():
    from catalogo import PRODUCTOS, clave_producto
    catalogo = {clave_producto(c, n) for f, c, n, _ in PRODUCTOS if f == 'cq-l1-tequenos'}
    casos = [('FULL QUESO', False, None), ('FULL QUESO', True, None), ('BOQUITAS', False, None),
             ('Tequeñón', False, None), ('Guayaba queso', False, None), ('Salmón', False, '10 unidades'),
             ('Salmón', False, '40 unidades'), ('Salmón', False, '50 unidades')]
    for nombre, freido, pres in casos:
        assert producto(nombre, None, pres, freido=freido)[0] in catalogo


def test_hay_freido_solo_con_mediciones():
    assert hay_freido({'tAceite': 174.0})
    assert hay_freido({'tInternaFreido': 52.0})
    assert not hay_freido({'tiempoFreido': 60, 'defectos': 3.0})


# ───────────── fusión y repetidos ─────────────
def test_fusionar_completa_huecos_sin_pisar_la_hoja_principal():
    principal = {'valores': {'tMasa': 33.6, 'aspectoFreido': None}, 'pendientes': {'aspectoFreido': 'Correcto'}}
    vieja = {'valores': {'tMasa': 34.0, 'sabor': 'Buena', 'aspectoFreido': 'Excelente'},
             'pendientes': {'color': 'Amarillento'}}
    r = fusionar(principal, vieja)
    assert r['valores'] == {'tMasa': 33.6, 'sabor': 'Buena', 'aspectoFreido': 'Excelente'}
    # lo pendiente de un campo que la otra hoja sí resolvió ya no se anota
    assert r['pendientes'] == {'color': 'Amarillento'}


def test_quitar_repetidos_conserva_la_primera_fila_del_lote():
    regs = [{'lote': '28861', 'fuente': 'f112'}, {'lote': '28862', 'fuente': 'f113'},
            {'lote': '28861', 'fuente': 'f114'}, {'lote': None, 'fuente': 'f115'}, {'lote': None, 'fuente': 'f116'}]
    quedan, repetidos = quitar_repetidos(regs)
    assert [r['fuente'] for r in quedan] == ['f112', 'f113', 'f115', 'f116']
    assert [(r['fuente'], primera['fuente']) for r, primera in repetidos] == [('f114', 'f112')]


# ───────────── extremo a extremo con un libro de prueba ─────────────
ENCABEZADO_CC = ['N°', 'FECHA DE PRODUCCIÓN ', 'PRODUCTO ', 'TEMPERATURA AMBIENTE (°C )', '% HUMEDAD AMBIENTE',
                 'TEMPERATURA (°C )\nDEL AGUA', 'TEMPERATURA (°C )\nMASA ', '% HUMEDAD MASA', 'PROVEEDOR',
                 '% HUMEDAD QUESO', 'APARIENCIA GENERAL\n', 'VUELTAS DE MASA', 'APARIENCIA GENERAL ',
                 'PESO PROMEDIO (g)', 'PESO TEÓRICO (g)', 'PESO PROMEDIO DEL RELLEN O (g)',
                 'PESO TEÓRICO DEL RELLENO (g)', 'TEMPERATURA INTERNA (°C)', 'TEMPERATURA \nACEITE (°C)',
                 'TIEMPO DE FREIDO', 'TEMPERATURA INTERNA  (°C) PCC1', 'ASPECTO GENERAL ',
                 'CANTIDAD TEQUEÑOS CON DEFECTO', 'TEMPERATURA DE PRODUCTO INICIAL\n ( °C )PCC2', 'APARIENCIA ',
                 'PRESENTACIÓN ', 'LOTE', 'FECHA DE VENCIMIENTO', 'T°C FINAL PCC3', 'PESO (kg) EMPAQUE',
                 'PESO REAL (kg)', 'DIFERENCIA (kg)', 'AUSENCIA/PRESENCIA DE METAL PCC4',
                 'TEMPERATURA DE SELLADO °C', 'CALIDAD DEL SELLADO', 'IMPRESIÓN DE LOTE']
ENCABEZADO_TQ = ['N°', 'FECHA DE PRODUCCIÓN ', 'PRODUCTO ', 'PRESENTACIÓN ', 'LOTE', 'FECHA DE VENCIMIENTO',
                 'PESO EMPAQUE', 'IMPRESIÓN DE LOTE', 'TEMPERATURA (°C )\nMASA ', '% HUMEDAD MASA',
                 'APARIENCIA GENERAL\n', 'VUELTAS DE MASA', 'APARIENCIA GENERAL ', 'PESO', 'TEMPERATURA \nACEITE (°C)',
                 'TIEMPO DE FREIDO', 'ASPECTO GENERAL ', 'CANTIDAD TEQUEÑOS CON DEFECTO',
                 'TEMPERATURA DE PRODUCTO\n ( °C )', 'APARIENCIA ', 'SABOR', 'COLOR', 'TEXTURA', 'APARIENCIA', 'FOTO']


def _fila_cc(n, fecha, prod, lote, vence, pres='24 UNIDADES', aceite=174, pcc1=52, agua='Ambiente', apariencia='Correcta',
             metal='Ausencia ', masa=31.9):
    return [n, fecha, prod, 29, 57, agua, masa, 30.91, 'quesos rafa', 52.78, apariencia, 5, 'Correcta', 37.09, '36-38',
            19.01, '16-18', -4, aceite, time(0, 1, 8), pcc1, 'Correcto', 4, 8.9, 'Correcta', pres, lote, vence, None,
            0.825, 0.64, 0.185, metal, 140, 'Correcta', 'Correcta']


@pytest.fixture
def carpeta(tmp_path):
    wb = openpyxl.Workbook()
    cc = wb.active
    cc.title = 'Control de calidad'
    cc.append(['LÍNEA 1-TEQUEÑOS'])
    cc.append([])
    cc.append([None] * 11 + ['PROCESO DE ENRROLLADO'])
    cc.append(ENCABEZADO_CC + [None, 'Fecha ', 'Temperatura del medidor exterior'])
    cc.append(_fila_cc(1, datetime(2025, 11, 12), 'FULL QUESO', '[PT026] 25132', datetime(2026, 11, 6), aceite='167.5°C',
                       pcc1='-', masa='33.5°C') + [None, datetime(2026, 11, 4), None])
    cc.append(_fila_cc(2, '23/04/2026', 'FULL QUESO', 'PT011(27217)', '23/10/2026', pres='12 UNIDADES', aceite='NA',
                       pcc1='NA', apariencia='Pálida'))
    # fecha ambigua: Excel la leyó como 5 de abril, pero por la fila anterior es el 4 de mayo
    cc.append(_fila_cc(3, datetime(2026, 4, 5), 'FULL QUESO', 'PT011(27217)', datetime(2026, 4, 11), metal='Presencia '))
    cc.append(_fila_cc(4, '15/04/2028', 'FULL QUESO', 'PT026 (27097)', '15/10/2026'))
    cc.append(_fila_cc(5, '25/04/2026', 'BOQUITAS', 'PT026(29609)', '25/10/2026', aceite='NA', pcc1='NA'))
    cc.append(_fila_cc(6, '27/04/2026', 'Mixto', '-', '-', agua=0.7))
    cc.append(['', None, None])
    tq = wb.create_sheet('TEQUEÑOS')
    tq.append([None, None, None, 'CALIDAD DEL EMPAQUE '])
    tq.append(ENCABEZADO_TQ)
    tq.append([2, datetime(2025, 11, 12), 'Tequeño full queso precocido', '24 unidades', '[PT026] 25132',
               datetime(2026, 11, 6), 1.009, 'Correcto', '33.5°C', 27.8, 'Adecuado', 4, 'Buena', 38, '167.5°C',
               '1 Minuto', 'Excelente', 4, '6.0°C', 'Buena', 'Insipido ', 'Amarillento', 'Buena', 'Regular', None])
    tq.append([3, '13/12/2025', None])
    wb.create_sheet('Control de masa').append(['Fecha: 25/05/2026'])
    wb.save(tmp_path / ARCHIVOS[0])
    return tmp_path


def test_extraer_de_punta_a_punta(carpeta):
    regs, rech = extraer(carpeta, HOY)
    por_lote = {r['lote']: r for r in regs}
    assert set(por_lote) == {'25132', '27217', '29609', None}
    assert len(regs) == 4

    r = por_lote['25132']
    assert r['formato'] == 'cq-l1-tequenos'
    assert r['fecha'] == '2025-12-11'
    assert r['hora'] is None and r['area'] is None and r['realizado_por'] is None
    assert r['producto'] == 'Tequeño Full Queso precocido'
    v = r['valores']
    assert v['tMasa'] == 33.5 and v['tAgua'] == 'Ambiente' and v['proveedorQueso'] == 'Quesos Rafa'
    assert v['tAceite'] == 167.5 and v['tiempoFreido'] == 68 and 'tInternaFreido' not in v
    assert v['vencimiento'] == '2026-06-11' and v['presentacion'] == '24 unidades'
    assert v['pesoEmpaque'] == 0.825 and v['metales'] == 'Ausencia' and v['sellado'] == 'Correcto'
    # de la hoja TEQUEÑOS: sensorial y aspecto del freído (en la principal decía 'Correcto', fuera de escala)
    assert v['aspectoFreido'] == 'Excelente' and v['textura'] == 'Buena' and v['aparienciaSensorial'] == 'Regular'
    assert 'color' not in v and 'sabor' not in v
    assert 'Color' in r['observaciones'] and 'Amarillento' in r['observaciones']
    assert 'Insipido' in r['observaciones']
    assert 'Correcto' not in r['observaciones']
    for calculado in ('pesoDeclarado', 'diferenciaPeso'):
        assert calculado not in v
    assert 'Control de calidad › fila 5' in r['fuente'] and 'TEQUEÑOS › fila 3' in r['fuente']

    crudo = por_lote['27217']
    assert crudo['producto'] == 'Tequeño Full Queso'
    assert not any(k in crudo['valores'] for k in ('tAceite', 'tiempoFreido', 'tInternaFreido', 'aspectoFreido',
                                                   'defectos'))
    assert 'aparienciaMasa' not in crudo['valores'] and 'Pálida' in crudo['observaciones']

    assert por_lote['29609']['producto'] is None
    sin_lote = por_lote[None]
    assert sin_lote['producto'] is None and sin_lote['valores']['tAgua'] == 0.7

    motivos = sorted(x['motivo'] for x in rech)
    assert motivos == ['fecha fuera de rango', 'lote repetido']
    repetido = next(x for x in rech if x['motivo'] == 'lote repetido')
    assert 'fila 7' in repetido['fuente'] and 'fila 6' in repetido['contenido']
    # el contenido muestra las fechas ya corregidas, no como las leyó Excel
    assert '04/05/2026' in repetido['contenido'] and '04/11/2026' in repetido['contenido']


def test_extraer_sin_la_hoja_principal_lo_reporta(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = 'Otra hoja'
    wb.save(tmp_path / ARCHIVOS[0])
    regs, rech = extraer(tmp_path, HOY)
    assert regs == []
    assert [x['motivo'] for x in rech] == ['hoja no encontrada']


def test_extraer_sin_el_archivo_lo_reporta(tmp_path):
    regs, rech = extraer(tmp_path, HOY)
    assert regs == []
    assert [x['motivo'] for x in rech] == ['archivo no encontrado']
