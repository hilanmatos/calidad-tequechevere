"""Extractor de 'CONTROL DE CALIDAD YUCA BITES.xlsx' → formato cq-l5-yuca (Línea 5 Yuca Bites y carimañola).

Hojas:
- 'Hoja1': la vigente. Encabezados en la fila 4, una fila por producción (lote) desde la fila 5.
- 'YUCA BITES': versión anterior de la hoja (dic-2025), con otro orden de columnas y tiempo/temperatura de
  cocción invertidos. Sus producciones ya están corregidas en Hoja1: cada fila se rechaza como duplicada (o,
  si no aparece en Hoja1, para revisión manual). Las filas que solo traen el N° de plantilla se ignoran.
- 'Hoja2': estudio de tiempos por carrito (sin fecha, producto ni lote): va a rechazos.

Uso: python -m extractores.yuca "<carpeta de Calidad>"   (desde tools/importar)
"""
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from extractores._bites import (aplicar_unicidad_lote, clave_lote, filas_con_datos, imprimir_resumen, lote_op,
                                procesar_filas, rechazo, unidades, verificar_encabezados)
from normalizar import corregir_secuencia

FORMATOS = ['cq-l5-yuca']
ARCHIVOS = ['CONTROL DE CALIDAD YUCA BITES.xlsx']

HOJA = 'Hoja1'
FILA_ENCABEZADOS = 4
HOJA_ANTERIOR = 'YUCA BITES'
FILA_ENCABEZADOS_ANTERIOR = 2
ENCABEZADOS_ANTERIOR = {2: 'FECHA DE PRODUCCIÓN', 4: 'PRESENTACIÓN', 5: 'LOTE'}
HOJA_TIEMPOS = 'Hoja2'
FILA_INICIO_TIEMPOS = 4

REGLAS = [
    (('yuca', 'queso'), 'Yuca Bite de Queso'),
    (('yuca', 'pollo'), 'Yuca Bite de Pollo'),
    (('yuca', 'carne'), 'Yuca Bite de Carne'),
    (('carimanola',), 'Mini Carimañola'),
]

CORRECTA = ('conf', ('Correcta', 'Incorrecta'))
CORRECTO = ('conf', ('Correcto', 'Incorrecto'))

ENCABEZADOS = {
    1: 'FECHA', 2: 'PRODUCTO', 3: '% HUMEDAD YUCA FRESCA', 4: 'TIEMPO DE COCCIÓN YUCA',
    5: 'TEMPERATURA DE COCCIÓN ( °C ) YUCA', 6: 'CANTIDAD DE AGUA', 7: 'TEMPERATURA (°C) MASA YUCA BITES',
    8: 'T° Ambiente proceso', 9: '%H Ambiente proceso', 10: '% HUMEDAD DE LA MASA YUCA BITES',
    11: 'APARIENCIA GENERAL', 12: 'TEMPERATURA DEL RELLENO (°C)', 13: 'Tiempo de espera', 14: 'T°C relleno final',
    15: 'TIEMPO DE FORMADO', 16: 'PESO PROMEDIO POR UNIDAD (g)', 17: 'PESO PROMEDIO DE RELLENO (g)',
    18: 'APARIENCIA', 19: '% DE PRODUCTO NO CONFORME', 20: 'TIEMPO', 21: 'TEMPERATURA INTERNA DEL PRODUCTO °C',
    22: 'TEMPERATURA ACEITE °C', 23: 'TIEMPO DE FRITURA', 24: 'TEMPERATURA INTERNA DE PRODUCTO ( °C ) PCC1',
    25: '% DE PRODUCTO NO CONFORME', 26: 'TEMPERATURA INTERNA DE PRODUCTO ( °C )', 27: 'APARIENCIA',
    28: 'PRESENTACIÓN', 29: 'LOTE', 30: 'FECHA DE VENCIMIENTO', 31: 'T°C INTERNA FINAL PCC2',
    32: 'PESO EMPAQUE (KG)', 33: 'PESO REAL', 34: 'DIFERENCIA DE PESO', 35: 'AUSENCIA/PRESENCIA DE METAL PCC3',
    36: 'CALIDAD DE SELLADO', 37: 'IMPRESIÓN DE LOTE', 38: 'SABOR', 39: 'COLOR', 40: 'TEXTURA', 41: 'APARIENCIA',
    42: 'FOTO',
}

# (columna, campo, tipo, etiqueta para observaciones). Col 1 = fecha, col 2 = producto.
COLUMNAS = [
    (3, 'hYuca', 'num', 'Humedad de la yuca fresca'),
    (4, 'tiempoCoccion', 'min', 'Tiempo de cocción de la yuca'),
    (5, 'tCoccion', 'num', 'Temperatura de cocción de la yuca'),
    (6, 'cantidadAgua', 'num', 'Cantidad de agua'),
    (7, 'tMasa', 'num', 'Temperatura de la masa'),
    (8, 'tAmbiente', 'num', 'Temperatura ambiente'),
    (9, 'hrAmbiente', 'num', 'Humedad relativa ambiente'),
    (10, 'hMasa', 'num', 'Humedad de la masa'),
    (11, 'aparienciaMasa', CORRECTA, 'Apariencia de la masa'),
    (12, 'tRelleno', 'num', 'Temperatura del relleno'),
    (13, 'tiempoEspera', 'min', 'Tiempo de espera del relleno'),
    (14, 'tRellenoFinal', 'num', 'Temperatura final del relleno'),
    (15, 'tiempoFormado', 'min', 'Tiempo de formado'),
    (16, 'pesoPromedio', 'num', 'Peso promedio por unidad'),
    (17, 'pesoRelleno', 'num', 'Peso promedio del relleno'),
    (18, 'aparienciaFormado', CORRECTA, 'Apariencia (formado)'),
    (19, 'pctNoConformeFormado', 'num', 'Producto no conforme (formado)'),
    (20, 'tiempoPreCong', 'min', 'Tiempo en congelación previa'),
    (21, 'tInternaPreCong', 'num', 'Temperatura interna (congelación previa)'),
    (22, 'tAceite', ('num', ('horno', 'aceite')), 'Temperatura del aceite u horno'),
    (23, 'tiempoFritura', 'seg', 'Tiempo de fritura'),
    (24, 'tInternaFritura', 'num', 'Temperatura interna (fritura)'),
    (25, 'pctNoConformeFritura', 'num', 'Producto no conforme (fritura)'),
    (26, 'tInternaCong', 'num', 'Temperatura interna (congelación)'),
    (27, 'aparienciaCong', CORRECTA, 'Apariencia (congelación)'),
    (28, 'presentacion', 'pres', 'Presentación'),
    (29, None, 'lote', 'Lote'),
    (30, 'vencimiento', 'venc', 'Fecha de vencimiento'),
    (31, 'tFinal', 'num', 'Temperatura interna final'),
    (32, 'pesoEmpaque', 'kg', 'Peso medido del empaque'),
    (33, None, 'ignorar', 'Peso real (es el declarado: lo calcula la app)'),
    (34, None, 'ignorar', 'Diferencia de peso (calculada por la app)'),
    (35, 'metales', 'metal', 'Detección de metales'),
    (36, 'sellado', CORRECTO, 'Calidad del sellado'),
    (37, 'impresionLote', CORRECTA, 'Impresión de lote'),
    (38, 'sabor', 'sens', 'Sabor (sensorial)'),
    (39, 'color', 'sens', 'Color (sensorial)'),
    (40, 'textura', 'sens', 'Textura (sensorial)'),
    (41, 'aparienciaSensorial', 'sens', 'Apariencia (sensorial)'),
    (42, None, 'obs', 'Foto'),
]


ANTERIOR_DUPLICADA = 'hoja anterior duplicada (la producción ya está en Hoja1)'
ANTERIOR_SIN_EQUIVALENTE = ('hoja anterior sin equivalente en Hoja1 (la hoja "YUCA BITES" es el formato anterior y '
                            'no se importa; revisar a mano)')
ESTUDIO_TIEMPOS = 'estudio de tiempos (Hoja2: por carrito, sin fecha, producto ni lote; no es un registro del formato)'


def _ref(fuente: str) -> str:
    return ' › '.join(fuente.split(' › ')[1:])


def rechazos_hoja_anterior(filas, archivo: str, registros: list[dict], hoy: date) -> list[dict]:
    """Cada fila de la hoja anterior se busca en Hoja1: por lote si lo tiene; si no, por fecha y unidades."""
    por_lote = {clave_lote(r['lote']): r for r in registros if r['lote']}
    por_fecha = {}
    for r in registros:
        por_fecha.setdefault((r['fecha'], unidades(r['valores'].get('presentacion'))), r)
    fechas = corregir_secuencia([f[1] if len(f) > 1 else None for _, f in filas], hoy)
    res = []
    for (n, fila), f in zip(filas, fechas):
        celda = lambda col: fila[col - 1] if col <= len(fila) else None  # noqa: E731
        lote, _ = lote_op(celda(5))
        if lote:
            igual = por_lote.get(clave_lote(lote))
        else:
            igual = por_fecha.get((f.isoformat(), unidades(celda(4)))) if f else None
        fuente = f'{archivo} › {HOJA_ANTERIOR} › fila {n}'
        if igual:
            res.append(rechazo(fuente, ANTERIOR_DUPLICADA, fila, f'misma producción en {_ref(igual["fuente"])}'))
        else:
            res.append(rechazo(fuente, ANTERIOR_SIN_EQUIVALENTE, fila))
    return res


def extraer(carpeta: Path, hoy: date) -> tuple[list[dict], list[dict]]:
    archivo = ARCHIVOS[0]
    wb = load_workbook(Path(carpeta) / archivo, read_only=True, data_only=True)
    try:
        conocidas = {HOJA, HOJA_ANTERIOR, HOJA_TIEMPOS}
        rechazos = [{'fuente': f'{archivo} › {ws.title}', 'motivo': 'hoja no reconocida (no se importó)',
                     'contenido': ''} for ws in wb.worksheets if ws.title not in conocidas]
        ws = wb[HOJA]
        verificar_encabezados(ws, FILA_ENCABEZADOS, ENCABEZADOS)
        filas = filas_con_datos(ws, FILA_ENCABEZADOS + 1)
        registros, rech = procesar_filas(filas, archivo, ws.title, FORMATOS[0], COLUMNAS, REGLAS, hoy)
        rechazos += rech
        if HOJA_ANTERIOR in wb.sheetnames:
            ant = wb[HOJA_ANTERIOR]
            verificar_encabezados(ant, FILA_ENCABEZADOS_ANTERIOR, ENCABEZADOS_ANTERIOR)
            filas_ant = filas_con_datos(ant, FILA_ENCABEZADOS_ANTERIOR + 1, ignorar_cols={1})
            rechazos += rechazos_hoja_anterior(filas_ant, archivo, registros, hoy)
        if HOJA_TIEMPOS in wb.sheetnames:
            rechazos += [rechazo(f'{archivo} › {HOJA_TIEMPOS} › fila {n}', ESTUDIO_TIEMPOS, fila)
                         for n, fila in filas_con_datos(wb[HOJA_TIEMPOS], FILA_INICIO_TIEMPOS)]
    finally:
        wb.close()
    registros, repetidos = aplicar_unicidad_lote(registros)
    return registros, rechazos + repetidos


if __name__ == '__main__':
    imprimir_resumen(*extraer(Path(sys.argv[1]), date.today()))
