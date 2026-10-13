"""Extractor de `temp-cuartos` (temperatura de cuartos de almacenamiento en frío).

El Excel tiene una sola hoja con una fila de título, una fila de encabezados y una fila por día:
Fecha + pistola y termómetro del cuarto de los cuatro cuartos (frío y congelación, planta baja y alta).
Cada fila es un registro (el formato admite uno por día). No hay hora ni responsable en el Excel.
"""
import json
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from normalizar import corregir_secuencia, numero, sin_acentos, texto, vacio

FORMATOS = ['temp-cuartos']
ARCHIVOS = ['Control de Temperatura-Cuartos de almacenamiento.xlsx']

DESDE = date(2025, 1, 1)
CUARTOS = {'frioPB': 'Cuarto frío PB', 'congPB': 'Congelación PB', 'frioPA': 'Cuarto frío PA', 'congPA': 'Congelación PA'}
MEDIDAS = {'Display': 'termómetro del cuarto', 'Pistola': 'pistola'}


def _norm(s) -> str:
    return re.sub(r'\s+', ' ', sin_acentos(str(s)).lower()).strip()


def campo_cuarto(encabezado):
    """Id del campo de la app para un encabezado del Excel ('Cuarto Frío PB Pistola' → 'frioPBPistola')."""
    if encabezado is None:
        return None
    s = _norm(encabezado)
    if 'planta baja' in s or re.search(r'pb(?![a-z])', s):
        planta = 'PB'
    elif 'planta alta' in s or re.search(r'pa(?![a-z])', s):
        planta = 'PA'
    else:
        return None
    if 'congel' in s:
        tipo = 'cong'
    elif 'frio' in s:
        tipo = 'frio'
    else:
        return None
    return f"{tipo}{planta}{'Pistola' if 'pistola' in s else 'Display'}"


def _etiqueta(campo: str) -> str:
    base, medida = re.match(r'(.+?)(Display|Pistola)$', campo).groups()
    return f'{CUARTOS[base]}: {MEDIDAS[medida]}'


def _corto(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    return str(v).strip()


def _contenido(fila) -> str:
    return ' | '.join(_corto(v) for v in fila if not vacio(v))[:160]


def registros_de_filas(filas, hoy: date, archivo: str, hoja: str):
    """Filas de una hoja (tuplas de celdas) → (registros, rechazos, encabezados sin destino)."""
    filas = [tuple(f) for f in filas]
    i_enc = next((i for i, f in enumerate(filas) if any(not vacio(v) and _norm(v) == 'fecha' for v in f)), None)
    if i_enc is None:
        primera = next((f for f in filas if any(not vacio(v) for v in f)), ())
        return [], [{'fuente': f'{archivo} › {hoja}', 'motivo': 'hoja sin encabezado reconocible',
                     'contenido': _contenido(primera)}], []

    col_fecha, col_obs, columnas, sin_destino = None, None, {}, []
    for j, enc in enumerate(filas[i_enc]):
        if vacio(enc):
            continue
        s = _norm(enc)
        if s == 'fecha':
            col_fecha = j
        elif s.startswith('observ'):
            col_obs = j
        elif campo_cuarto(enc):
            columnas[j] = campo_cuarto(enc)
        else:
            sin_destino.append(str(enc).strip())

    datos = [(i + 1, f) for i, f in enumerate(filas) if i > i_enc and any(not vacio(v) for v in f)]
    celda = lambda f, j: f[j] if j is not None and j < len(f) else None
    fechas = corregir_secuencia([celda(f, col_fecha) for _, f in datos], hoy)

    registros, rechazos, vistos = [], [], {}
    for (n, f), d in zip(datos, fechas):
        fuente = f'{archivo} › {hoja} › fila {n}'
        rechazo = lambda motivo: rechazos.append({'fuente': fuente, 'motivo': motivo, 'contenido': _contenido(f)})
        if d is None:
            rechazo('sin fecha' if vacio(celda(f, col_fecha)) else 'fecha ilegible')
            continue
        if not DESDE <= d <= hoy:
            rechazo('fecha fuera de rango (2025-01-01 a hoy)')
            continue
        valores, avisos = {}, []
        for j, campo in columnas.items():
            v = celda(f, j)
            if vacio(v):
                continue
            x = numero(v)
            if x is None:
                avisos.append(f"{_etiqueta(campo)} sin número en el Excel: '{_corto(v)}'")
            else:
                valores[campo] = x
        if not valores:
            rechazo('sin lecturas de temperatura')
            continue
        if d in vistos:
            rechazo(f'día repetido: ya hay un registro de {d.isoformat()} (fila {vistos[d]})')
            continue
        vistos[d] = n
        obs = [o for o in [texto(celda(f, col_obs))] + avisos if o]
        registros.append({
            'formato': FORMATOS[0],
            'fecha': d.isoformat(),
            'hora': None,
            'area': None,
            'producto': None,
            'lote': None,
            'valores': valores,
            'observaciones': '; '.join(obs) or None,
            'accion_correctiva': None,
            'realizado_por': None,
            'fuente': fuente,
        })
    return registros, rechazos, sin_destino


def _extraer(carpeta: Path, hoy: date):
    registros, rechazos, sin_destino = [], [], []
    for archivo in ARCHIVOS:
        ruta = Path(carpeta) / archivo
        if not ruta.exists():
            rechazos.append({'fuente': archivo, 'motivo': 'archivo no encontrado', 'contenido': str(ruta)})
            continue
        wb = load_workbook(ruta, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                r, x, s = registros_de_filas(ws.iter_rows(values_only=True), hoy, archivo, ws.title)
                registros += r
                rechazos += x
                sin_destino += [f'{ws.title} › {c}' for c in s]
        finally:
            wb.close()
    return registros, rechazos, sin_destino


def extraer(carpeta: Path, hoy: date):
    registros, rechazos, _ = _extraer(carpeta, hoy)
    return registros, rechazos


def imprimir_resumen(registros, rechazos, sin_destino):
    print('Registros por formato:')
    for formato, n in sorted(Counter(r['formato'] for r in registros).items()):
        print(f'  {formato}: {n}')
    if registros:
        fechas = sorted(r['fecha'] for r in registros)
        print(f'Rango de fechas: {fechas[0]} a {fechas[-1]}')
    print(f'Rechazos: {len(rechazos)}')
    for motivo, n in Counter(r['motivo'] for r in rechazos).most_common():
        print(f'  {n} × {motivo}')
    if sin_destino:
        print('Columnas sin destino:', ', '.join(sin_destino))
    print('Muestra:')
    for r in registros[:2]:
        print(json.dumps(r, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    if len(sys.argv) != 2:
        sys.exit('uso: python -m extractores.temp_cuartos "<carpeta de Calidad>"')
    imprimir_resumen(*_extraer(Path(sys.argv[1]), date.today()))
