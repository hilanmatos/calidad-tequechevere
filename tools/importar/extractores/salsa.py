"""Extractor de `cq-salsa-tartara` (control de calidad de la Salsa tártara).

El Excel tiene una hoja con título, una fila de PCC, una fila de secciones y una fila de encabezados; debajo,
una fila por producción. Cada fila es un registro (el formato admite uno por lote). Las columnas se ubican por
el texto del encabezado, no por su posición. No hay hora, área ni responsable en el Excel.

Fechas: las de Excel con día ≤ 12 vienen con día y mes invertidos (se escribieron dd/mm en un Excel m/d).
La de producción se resuelve con normalizar.corregir_secuencia; la de vencimiento, contra la de producción
(vencimiento_corregido).
"""
import json
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from normalizar import conforme, corregir_secuencia, fecha, numero, sin_acentos, texto, vacio

FORMATOS = ['cq-salsa-tartara']
ARCHIVOS = ['CONTROL DE CALIDAD SALSA TÁRTARA.xlsx']

DESDE = date(2025, 1, 1)

# Encabezado normalizado (minúsculas, sin acentos) → destino. Gana el primer patrón que coincide.
# Destinos con '_' no van a `valores`: _n (correlativo), _teorico y _diferencia (los calcula la app).
COLUMNAS = [
    (r'^(n|no|nro)\W*$', '_n'),
    (r'^fecha de produccion', 'fecha'),
    (r'^producto$', 'producto'),
    (r'^temperatura ambiente', 'tAmbiente'),
    (r'humedad', 'hrAmbiente'),
    (r'^apariencia vegetales', 'aparienciaVegetales'),
    (r'^lavado', 'lavado'),
    (r'^sanitec', 'sanitec'),
    (r'^tiempo higienizacion', 'tiempoHigienizacion'),
    (r'^enjuague', 'enjuague'),
    (r'^escurrido', 'escurrido'),
    (r'^despepitado|corte', 'corte'),
    (r'^picado', 'picado'),
    (r'^peso vegetal', 'pesoVegetal'),
    (r'^peso mayonesa', 'pesoMayonesa'),
    (r'^peso vinagre', 'pesoVinagre'),
    (r'^peso sal\b', 'pesoSal'),
    (r'^tiempo mezclado', 'tiempoMezclado'),
    (r'^apariencia mezcla', 'aparienciaMezcla'),
    (r'^ph\b', 'ph'),
    (r'^color', 'colorSalsa'),
    (r'^olor', 'olor'),
    (r'^consistencia', 'consistencia'),
    (r'envasado', 'tEnvasado'),
    (r'^tiempo expuesto', 'tiempoExpuesto'),
    (r'^peso neto', 'pesoNeto'),
    (r'^peso teorico', '_teorico'),
    (r'^diferencia', '_diferencia'),
    (r'^etiquetado', 'etiquetado'),
    (r'^lote', 'lote'),
    (r'^fecha de vencimiento', 'vencimiento'),
    (r'^tapado|cierre', 'cierre'),
    (r'metales', 'metales'),
    (r'^inspeccion visual', 'inspeccionVisual'),
    (r'almacenamiento', 'tAlmacenamiento'),
    (r'^observ', 'observaciones'),
]

NUMEROS = {'tAmbiente', 'hrAmbiente', 'sanitec', 'tiempoHigienizacion', 'tiempoMezclado', 'ph', 'tEnvasado',
           'tiempoExpuesto', 'pesoNeto', 'tAlmacenamiento'}
CONFORMES = {c: ('Conforme', 'No conforme') for c in (
    'lavado', 'enjuague', 'escurrido', 'corte', 'picado', 'pesoVegetal', 'pesoMayonesa', 'pesoVinagre', 'pesoSal',
    'etiquetado', 'inspeccionVisual')}
CONFORMES.update(aparienciaVegetales=('Correcta', 'Incorrecta'), cierre=('Correcto', 'Incorrecto'))
OPCIONES = {
    'aparienciaMezcla': ['Homogénea', 'No homogénea'],
    'colorSalsa': ['Característico', 'Atípico'],
    'olor': ['Característico', 'Atípico'],
    'consistencia': ['Cremosa y homogénea', 'Líquida', 'Grumosa'],
    'metales': ['Ausencia', 'Presencia'],
}
ETIQUETAS = {
    'tAmbiente': 'Temperatura ambiente', 'hrAmbiente': 'Humedad relativa', 'aparienciaVegetales': 'Apariencia de los vegetales',
    'lavado': 'Lavado', 'sanitec': 'Sanitec', 'tiempoHigienizacion': 'Tiempo de higienización', 'enjuague': 'Enjuague',
    'escurrido': 'Escurrido', 'corte': 'Despepitado y corte', 'picado': 'Picado', 'pesoVegetal': 'Peso del vegetal',
    'pesoMayonesa': 'Peso de la mayonesa', 'pesoVinagre': 'Peso del vinagre', 'pesoSal': 'Peso de la sal',
    'tiempoMezclado': 'Tiempo de mezclado', 'aparienciaMezcla': 'Apariencia de la mezcla', 'ph': 'pH',
    'colorSalsa': 'Color', 'olor': 'Olor', 'consistencia': 'Consistencia', 'tEnvasado': 'Temperatura al envasar',
    'tiempoExpuesto': 'Tiempo expuesto', 'pesoNeto': 'Peso neto', 'etiquetado': 'Etiquetado',
    'vencimiento': 'Fecha de vencimiento', 'cierre': 'Tapado y cierre', 'metales': 'Detección de metales',
    'inspeccionVisual': 'Inspección visual', 'tAlmacenamiento': 'Temperatura de almacenamiento',
}


def _norm(s) -> str:
    return re.sub(r'\s+', ' ', sin_acentos(str(s)).lower()).strip()


def _destino(encabezado):
    s = _norm(encabezado)
    return next((d for patron, d in COLUMNAS if re.search(patron, s)), None)


def producto_salsa(nombre, teorico=None):
    """Código del catálogo: PS002 (600 g) o PS001 (galón de 4,000 g); None si el Excel no permite distinguirlo."""
    if not vacio(nombre):
        s = re.sub(r'(?<=\d)[.,](?=\d{3})', '', _norm(nombre))
        if 'galon' in s or re.search(r'(?<!\d)4000(?!\d)', s):
            return 'PS001'
        if re.search(r'(?<!\d)600(?!\d)', s):
            return 'PS002'
    return {600: 'PS002', 4000: 'PS001'}.get(numero(teorico))


def _tokens(s) -> frozenset:
    return frozenset(t for t in re.findall(r'[a-z0-9]+', _norm(s)) if t != 'y')


def opcion_de(v, opciones):
    """La opción de la app con las mismas palabras que la celda ('Cremosa / homogénea' → 'Cremosa y homogénea')."""
    if vacio(v):
        return None
    t = _tokens(v)
    return next((o for o in opciones if _tokens(o) == t), None)


def vencimiento_corregido(v, produccion: date):
    """Fecha de vencimiento con día y mes corregidos usando la fecha de producción ya resuelta.

    Texto dd/mm/aaaa y fechas de Excel con día > 12 son exactas. Si es ambigua, se descarta la versión que no
    queda después de la producción y, entre dos posibles, se elige la que cae el mismo día del mes que la
    producción (el vencimiento se fija sumando meses). Sin pista, queda como la guardó Excel.
    """
    d = fecha(v)
    if d is None or isinstance(v, str) or d.day > 12 or d.day == d.month:
        return d
    invertida = date(d.year, d.day, d.month)
    candidatos = [c for c in (d, invertida) if c > produccion]
    if len(candidatos) == 1:
        return candidatos[0]
    mismo_dia = [c for c in candidatos if c.day == produccion.day]
    return mismo_dia[0] if len(mismo_dia) == 1 else d


def lote_texto(v):
    if vacio(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _corto(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    return str(v).strip()


def _contenido(fila) -> str:
    return ' | '.join(_corto(v) for v in fila if not vacio(v))[:160]


def _valor(campo, v):
    if campo in NUMEROS:
        return numero(v)
    if campo in CONFORMES:
        return conforme(v, CONFORMES[campo])
    return opcion_de(v, OPCIONES[campo])


def registros_de_filas(filas, hoy: date, archivo: str, hoja: str):
    """Filas de una hoja (tuplas de celdas) → (registros, rechazos, encabezados sin destino)."""
    filas = [tuple(f) for f in filas]
    i_enc = next((i for i, f in enumerate(filas)
                  if any(not vacio(v) and _destino(v) == 'fecha' for v in f)), None)
    if i_enc is None:
        primera = next((f for f in filas if any(not vacio(v) for v in f)), ())
        return [], [{'fuente': f'{archivo} › {hoja}', 'motivo': 'hoja sin encabezado reconocible',
                     'contenido': _contenido(primera)}], []

    col, sin_destino = {}, []
    for j, enc in enumerate(filas[i_enc]):
        if vacio(enc):
            continue
        d = _destino(enc)
        if d is None or d in col:
            sin_destino.append(str(enc).strip())
        else:
            col[d] = j

    celda = lambda f, campo: f[col[campo]] if campo in col and col[campo] < len(f) else None
    con_datos = lambda f: any(not vacio(v) for j, v in enumerate(f) if j != col.get('_n'))
    datos = [(i + 1, f) for i, f in enumerate(filas) if i > i_enc and con_datos(f)]
    fechas = corregir_secuencia([celda(f, 'fecha') for _, f in datos], hoy)

    registros, rechazos, lotes = [], [], {}
    for (n, f), d in zip(datos, fechas):
        fuente = f'{archivo} › {hoja} › fila {n}'
        rechazo = lambda motivo: rechazos.append({'fuente': fuente, 'motivo': motivo, 'contenido': _contenido(f)})
        if d is None:
            rechazo('sin fecha' if vacio(celda(f, 'fecha')) else 'fecha ilegible')
            continue
        if not DESDE <= d <= hoy:
            rechazo('fecha fuera de rango (2025-01-01 a hoy)')
            continue
        lote = lote_texto(celda(f, 'lote'))
        clave = re.sub(r'\s+', '', lote).upper() if lote else None
        if clave in lotes:
            rechazo(f'lote repetido: {lote} ya está en la fila {lotes[clave]}')
            continue

        valores, avisos = {}, []
        for campo in col:
            if campo.startswith('_') or campo in ('fecha', 'producto', 'lote', 'observaciones'):
                continue
            v = celda(f, campo)
            if vacio(v):
                continue
            x = vencimiento_corregido(v, d) if campo == 'vencimiento' else _valor(campo, v)
            if x is None:
                avisos.append(f"{ETIQUETAS[campo]} sin valor reconocible en el Excel: '{_corto(v)}'")
            else:
                valores[campo] = x.isoformat() if isinstance(x, date) else x

        producto = producto_salsa(celda(f, 'producto'), celda(f, '_teorico'))
        if producto is None and not vacio(celda(f, 'producto')):
            avisos.insert(0, f"Producto en el Excel: '{_corto(celda(f, 'producto'))}'")
        if clave:
            lotes[clave] = n
        obs = [o for o in [texto(celda(f, 'observaciones'))] + avisos if o]
        registros.append({
            'formato': FORMATOS[0],
            'fecha': d.isoformat(),
            'hora': None,
            'area': None,
            'producto': producto,
            'lote': lote,
            'valores': valores,
            'observaciones': '; '.join(obs) or None,
            'accion_correctiva': None,
            'realizado_por': None,
            'fuente': fuente,
        })
    return registros, rechazos, sin_destino


def _extraer(carpeta: Path, hoy: date):
    registros, rechazos, sin_destino = [], [], []
    for archivo in ARCHIVOS:
        ruta = Path(carpeta) / archivo
        if not ruta.exists():
            rechazos.append({'fuente': archivo, 'motivo': 'archivo no encontrado', 'contenido': str(ruta)})
            continue
        wb = load_workbook(ruta, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                r, x, s = registros_de_filas(ws.iter_rows(values_only=True), hoy, archivo, ws.title)
                registros += r
                rechazos += x
                sin_destino += [f'{ws.title} › {c}' for c in s]
        finally:
            wb.close()
    return registros, rechazos, sin_destino


def extraer(carpeta: Path, hoy: date):
    registros, rechazos, _ = _extraer(carpeta, hoy)
    return registros, rechazos


def imprimir_resumen(registros, rechazos, sin_destino):
    print('Registros por formato:')
    for formato, n in sorted(Counter(r['formato'] for r in registros).items()):
        print(f'  {formato}: {n}')
    if registros:
        fechas = sorted(r['fecha'] for r in registros)
        print(f'Rango de fechas: {fechas[0]} a {fechas[-1]}')
    print('Productos:', dict(Counter(r['producto'] for r in registros)))
    print(f'Rechazos: {len(rechazos)}')
    for motivo, n in Counter(r['motivo'] for r in rechazos).most_common():
        print(f'  {n} × {motivo}')
    if sin_destino:
        print('Columnas sin destino:', ', '.join(sin_destino))
    print('Muestra:')
    for r in registros[:2]:
        print(json.dumps(r, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    if len(sys.argv) != 2:
        sys.exit('uso: python -m extractores.salsa "<carpeta de Calidad>"')
    imprimir_resumen(*_extraer(Path(sys.argv[1]), date.today()))
