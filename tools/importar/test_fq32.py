from datetime import date

import pytest
from docx import Document
from docx.oxml import parse_xml

from extractores.fq32 import (ARCHIVOS, apariencia, columnas_extra, condiciones, extraer, fecha_hora, hora_texto,
                              leer_fila, pagina_a_registro, paginas, peso_g, producto_catalogo, texto_celda)

HOY = date(2026, 10, 12)
TQ = ['cq-l1-tequenos']
EM = ['cq-l2-empanadas-trigo', 'cq-l3-empanadas-maiz']
GF = ['cq-l4-almond', 'cq-l5-yuca']
INSTRUCCION = ('Instrucción: tomar muestras de cada una de las mesas de trabajo, registrar nombre del operador, '
               'nombre del producto, peso en gramos, apariencia del producto marcar con un (√) en la casilla '
               '"C" si cumple o (X) en la casilla "NC" no cumple.  Fecha: {f}     Hora: {h}')


# --- fecha_hora: lo escrito tras "Fecha:" y "Hora:" en la instrucción de cada página

def test_fecha_hora_de_plantilla_en_blanco_es_none():
    assert fecha_hora(INSTRUCCION.format(f='___________________', h='_____________')) == (None, None)


def test_fecha_hora_lee_lo_escrito_sobre_la_raya():
    assert fecha_hora(INSTRUCCION.format(f='15/09/2026', h='8:30 am')) == ('15/09/2026', '8:30 am')


def test_fecha_hora_quita_rayas_sobrantes():
    assert fecha_hora(INSTRUCCION.format(f='__15/09/2026____', h='____')) == ('15/09/2026', None)


def test_fecha_hora_sin_rotulos_es_none():
    assert fecha_hora('Pastelitos') == (None, None)


# --- peso_g: las muestras se anotan en gramos

@pytest.mark.parametrize('v, esperado', [('36', 36.0), ('36,5', 36.5), ('36 g', 36.0), ('0.036 kg', 36.0),
                                         ('', None), ('-', None), ('x', None)])
def test_peso_g(v, esperado):
    assert peso_g(v) == esperado


# --- apariencia: marca en la casilla C o en la casilla NC

def test_apariencia_marca_en_c():
    assert apariencia('√', '') == 'C'


def test_apariencia_marca_en_nc():
    assert apariencia('', 'X') == 'NC'


def test_apariencia_sin_marca_es_none():
    assert apariencia(' ', '') is None


def test_apariencia_con_las_dos_marcas_es_ambigua():
    assert apariencia('√', 'X') is None


# --- columnas_extra: las dos últimas columnas cambian según el documento

def test_columnas_extra_tequenos_y_gluten_free():
    assert columnas_extra(['Observaciones', 'Realizado por:']) == ['observacion', 'realizado']


def test_columnas_extra_empanadas():
    assert columnas_extra(['Lote', 'Fecha de vencimiento']) == ['lote', 'vencimiento']


def test_columnas_extra_desconocida_se_ignora():
    assert columnas_extra(['Comentario', 'Lote']) == [None, 'lote']


# --- leer_fila: una mesa de trabajo (17 celdas: N°, operador, producto, M1..M10, C, NC, extra, extra)

def fila(operador='', producto='', muestras=(), c='', nc='', extra1='', extra2='', n='1'):
    m = list(muestras) + [''] * (10 - len(muestras))
    return [n, operador, producto, *m, c, nc, extra1, extra2]


def test_leer_fila_en_blanco_es_none():
    assert leer_fila(fila(), ['observacion', 'realizado']) is None


def test_leer_fila_con_datos():
    r = leer_fila(fila('Olga', 'Boquita', ['27', '28,5', '', '29'], c='√', extra1='Masa blanda',
                       extra2='Alicia'), ['observacion', 'realizado'])
    assert r == {
        'operador': 'Olga', 'producto': 'Boquita',
        'muestras': [27.0, 28.5, None, 29.0, None, None, None, None, None, None],
        'apariencia': 'C', 'observacion': 'Masa blanda', 'realizado': 'Alicia', 'avisos': [],
    }


def test_leer_fila_empanadas_trae_lote_y_vencimiento():
    r = leer_fila(fila('Hermel', 'Pastelito de queso', ['30'], extra1='2609', extra2='15/03/2027'),
                  ['lote', 'vencimiento'])
    assert r['lote'] == '2609'
    assert r['vencimiento'] == '15/03/2027'
    assert r['apariencia'] is None


def test_leer_fila_avisa_si_marcaron_c_y_nc():
    r = leer_fila(fila('Olga', 'Boquita', ['27'], c='√', nc='X'), ['observacion', 'realizado'])
    assert r['apariencia'] is None
    assert r['avisos'] == ['apariencia marcada en C y en NC']


# --- producto_catalogo: texto escrito → clave del catálogo (código o nombre exacto) de las líneas del formato

@pytest.mark.parametrize('texto, lineas, esperado', [
    ('Tequeño Boquita', TQ, 'Tequeño Boquita'),
    ('boquita', TQ, 'Tequeño Boquita'),
    ('Tequeños full queso', TQ, 'Tequeño Full Queso'),
    ('Full 12', TQ, 'Tequeño Full Queso'),
    ('Tequeñón', TQ, 'Tequeñón'),
    ('Salmón 10', TQ, 'PT035'),
    ('Pastelito de queso', EM, 'PP001'),
    ('Empanadita maiz queso', EM, 'PE006'),
    ('Yuca bites pollo', GF, 'Yuca Bite de Pollo'),
    ('Mini carimañola', GF, 'Mini Carimañola'),
])
def test_producto_catalogo_reconoce(texto, lineas, esperado):
    assert producto_catalogo(texto, lineas) == esperado


@pytest.mark.parametrize('texto, lineas', [
    ('Salmón', TQ),          # tres presentaciones con códigos distintos
    ('Tequeño', TQ),         # demasiado general
    ('Queso', GF),           # Almond Bite y Yuca Bite de queso
    ('Pastelito de queso', TQ),  # existe, pero no en las líneas de este formato
    ('Almon Bites Queso', GF),   # mal escrito: no se adivina
    ('de', TQ),                  # solo palabras de relleno
    ('', TQ),
    (None, TQ),
])
def test_producto_catalogo_no_adivina(texto, lineas):
    assert producto_catalogo(texto, lineas) is None


# --- condiciones: tablas etiqueta | valor (de a pares por fila) → campos del formato

def test_condiciones_tequenos():
    filas = [
        ['Ti°Ccongelación', '-18', 'Tf°congelación', '-20'],
        ['T°C fritura', '175', 'T°C Aceite', '180,5'],
        ['Ti° congelación (Empaque)', '-19', 'Tf° congelación (Empaque)', ''],
        ['T°C Ambiente', '27', '%H Ambiente', '65'],
        ['T°C Masa', '22', '%H Masa', '35,2'],
        ['Lote queso', 'Q-2609', '%H Queso', '48'],
        ['Lote prod.', '260915', 'N°bandejas/ciclo', '12'],
    ]
    valores, conflictos = condiciones(filas)
    assert valores == {
        'tiCongelacion': -18.0, 'tfCongelacion': -20.0, 'tFritura': 175.0, 'tAceite': 180.5,
        'tiCongEmpaque': -19.0, 'tAmbiente': 27.0, 'hrAmbiente': 65.0, 'tMasa': 22.0, 'hMasa': 35.2,
        'loteQueso': 'Q-2609', 'hQueso': 48.0, 'loteProducto': '260915', 'bandejasCiclo': 12.0,
    }
    assert conflictos == []


def test_condiciones_empanadas_peso_de_empaque_en_kg():
    valores, _ = condiciones([['%H de la masa', '34'], ['T masa', '21'], ['T relleno', '4'],
                              ['W de Empaque', '484 g']])
    assert valores == {'hMasa': 34.0, 'tMasa': 21.0, 'tRelleno': 4.0, 'pesoEmpaque': 0.484}


def test_condiciones_plantilla_vacia_no_trae_valores():
    assert condiciones([['T°C Ambiente', '', '%H Ambiente', ''], ['T masa', '']]) == ({}, [])


def test_condiciones_ignora_etiquetas_desconocidas():
    assert condiciones([['Almond Bites Chocolate', '12.5', '2.50', '1.0']]) == ({}, [])


def test_condiciones_dos_valores_distintos_no_se_eligen():
    # Empanadas trae dos tablas iguales (pastelitos y empanaditas) para un solo juego de campos en la app
    valores, conflictos = condiciones([['T masa', '21'], ['%H de la masa', '34'],
                                       ['T masa', '23'], ['%H de la masa', '34']])
    assert valores == {'hMasa': 34.0}
    assert conflictos == ['T masa: 21 / 23']


# --- hora_texto: la hora se escribe a mano, a veces con am/pm

@pytest.mark.parametrize('v, esperado', [('8:30', '08:30'), ('8:30 am', '08:30'), ('2:15 pm', '14:15'),
                                         ('2:15 p.m.', '14:15'), ('12:10 pm', '12:10'), ('12:05 am', '00:05'),
                                         ('14:00', '14:00'), (None, None), ('mañana', None)])
def test_hora_texto(v, esperado):
    assert hora_texto(v) == esperado


# --- pagina_a_registro: una página del Word ya leída → registro, o rechazo

ARCHIVO_TQ, ARCHIVO_EM, ARCHIVO_GF = ARCHIVOS
ENC_TQ = [['N°', 'Operador', 'Producto', 'Peso (g) Muestra evaluada', 'Apariencia', 'Observaciones', 'Realizado por:'],
          ['', '', '', 'M1', 'M2', 'M3', 'M4', 'M5', 'M6', 'M7', 'M8', 'M9', 'M10', 'C', 'NC', '', '']]
ENC_EM = [['N°', 'Operador', 'Producto', 'Peso (g) Muestra evaluada', 'Apariencia', 'Lote', 'Fecha de vencimiento'],
          ENC_TQ[1]]
COND_TQ_VACIA = [['Ti°Ccongelación', '', 'Tf°congelación', ''], ['T°C Ambiente', '', '%H Ambiente', '']]
BLANCAS = [fila(n=str(i)) for i in range(1, 9)]


def pagina(f='___________________', h='_____________', enc=ENC_TQ, filas=BLANCAS, cond=COND_TQ_VACIA,
           pie='Revisado por: ______________________________', n=1):
    return {'n': n, 'instruccion': INSTRUCCION.format(f=f, h=h), 'muestras': enc + filas, 'condiciones': cond,
            'pie': pie}


def test_pagina_en_blanco_no_da_registro_ni_rechazo():
    assert pagina_a_registro(pagina(), ARCHIVO_TQ, HOY) == (None, [])


def test_pagina_tequenos_llena():
    filas = [fila('Olga', 'Boquita', ['27', '28', '27,5'], c='√', extra2='Alicia', n='1'),
             fila('moises', 'Tequeño Full Queso', ['37'], nc='X', extra1='Abiertos', extra2='Alicia', n='2')]
    cond = [['Ti°Ccongelación', '-18', 'Tf°congelación', ''], ['T°C Ambiente', '27', '%H Ambiente', '']]
    reg, rechazos = pagina_a_registro(pagina('15/09/2026', '2:30 pm', filas=filas + BLANCAS[2:], cond=cond, n=2),
                                      ARCHIVO_TQ, HOY)
    assert rechazos == []
    assert reg == {
        'formato': 'fq32-tequenos', 'fecha': '2026-09-15', 'hora': '14:30', 'area': 'Planta Alta',
        'producto': None, 'lote': None,
        'valores': {
            'muestras': [
                {'operario': 'Olga Palacios', 'producto': 'Tequeño Boquita',
                 'muestras': [27.0, 28.0, 27.5] + [None] * 7, 'apariencia': 'C'},
                {'operario': 'Moisés Silva', 'producto': 'Tequeño Full Queso',
                 'muestras': [37.0] + [None] * 9, 'apariencia': 'NC', 'observacion': 'Abiertos'},
            ],
            'tiCongelacion': -18.0, 'tAmbiente': 27.0,
        },
        'observaciones': None, 'accion_correctiva': None, 'realizado_por': 'Alicia',
        'fuente': f'{ARCHIVO_TQ} › página 2',
    }


def test_pagina_gluten_free_va_al_area_gluten_free():
    reg, _ = pagina_a_registro(pagina('01/09/2026', filas=[fila('Ricardo', 'Mini carimañola', ['32'])]),
                               ARCHIVO_GF, HOY)
    assert (reg['formato'], reg['area']) == ('fq32-gluten-free', 'Gluten Free')
    assert reg['valores']['muestras'][0]['producto'] == 'Mini Carimañola'
    assert reg['hora'] is None


def test_pagina_empanadas_lote_vencimiento_y_realizado_en_el_pie():
    filas = [fila('Hermel', 'Pastelito de queso', ['30', '31'], c='√', extra1='2609', extra2='15/03/2027')]
    cond = [['%H de la masa', '34'], ['T masa', '21'], ['%H de la masa', '35'], ['T masa', '']]
    reg, _ = pagina_a_registro(pagina('20/09/2026', '9:00', enc=ENC_EM, filas=filas, cond=cond,
                                      pie='  Realizado por: Alicia Sánchez \t\tRevisado por: ________'),
                               ARCHIVO_EM, HOY)
    assert reg['formato'] == 'fq32-empanadas'
    assert reg['valores'] == {
        'muestras': [{'operario': 'Hermel', 'producto': 'PP001', 'muestras': [30.0, 31.0] + [None] * 8,
                      'apariencia': 'C', 'lote': '2609', 'vencimiento': '2027-03-15'}],
        'tMasa': 21.0,
    }
    assert reg['observaciones'] == 'Condiciones con dos valores distintos: %H de la masa: 34 / 35'
    assert reg['realizado_por'] == 'Alicia Sánchez'


def test_pagina_anota_lo_que_no_se_pudo_pasar():
    filas = [fila('Olga', 'Almon queso', ['16'], c='√', nc='X', n='3')]
    reg, _ = pagina_a_registro(pagina('01/09/2026', filas=filas), ARCHIVO_GF, HOY)
    assert reg['valores']['muestras'][0]['producto'] is None
    assert 'apariencia' not in reg['valores']['muestras'][0]
    assert reg['observaciones'] == ("Mesa 3: producto 'Almon queso' no está en el catálogo; "
                                    'Mesa 3: apariencia marcada en C y en NC')


def test_pagina_vencimiento_ilegible_se_anota():
    filas = [fila('Hermel', 'PP001', ['30'], extra2='marzo')]
    reg, _ = pagina_a_registro(pagina('20/09/2026', enc=ENC_EM, filas=filas, cond=[]), ARCHIVO_EM, HOY)
    assert 'vencimiento' not in reg['valores']['muestras'][0]
    assert "Mesa 1: vencimiento ilegible 'marzo'" in reg['observaciones']


@pytest.mark.parametrize('f, motivo', [('___', 'sin fecha'), ('sept 15', 'fecha ilegible'),
                                       ('15/09/2024', 'fecha fuera de rango'), ('13/10/2026', 'fecha fuera de rango')])
def test_pagina_con_datos_y_fecha_mala_se_rechaza(f, motivo):
    reg, rechazos = pagina_a_registro(pagina(f, filas=[fila('Olga', 'Boquita', ['27'])], n=2), ARCHIVO_TQ, HOY)
    assert reg is None
    assert [(r['fuente'], r['motivo']) for r in rechazos] == [(f'{ARCHIVO_TQ} › página 2', motivo)]
    assert rechazos[0]['contenido']


# --- lectura del Word

NS_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def test_texto_celda_ignora_cuadros_de_texto_anclados():
    # en empanadas, el rótulo "Empanaditas" es un cuadro de texto anclado dentro de una celda de valor
    tc = parse_xml(f'<w:tc {NS_W}><w:p><w:r><w:t>34</w:t></w:r><w:r><w:pict><w:txbxContent><w:p><w:r>'
                   f'<w:t>Empanaditas</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p></w:tc>')
    assert texto_celda(tc) == '34'


def test_texto_celda_une_parrafos_y_cuenta_simbolos_como_marca():
    tc = parse_xml(f'<w:tc {NS_W}><w:p><w:r><w:t>Realizado</w:t></w:r></w:p><w:p><w:r><w:t>por:</w:t></w:r></w:p>'
                   f'<w:p><w:r><w:sym w:font="Symbol" w:char="F0D6"/></w:r></w:p></w:tc>')
    assert texto_celda(tc) == 'Realizado por: √'


def _tabla(doc, filas):
    t = doc.add_table(rows=len(filas), cols=max(len(f) for f in filas))
    for fila_t, valores in zip(t.rows, filas):
        for celda, v in zip(fila_t.cells, valores):
            celda.text = v
    return t


def docx_fq32(ruta, por_pagina):
    """Word con la estructura de la plantilla: por página, instrucción + muestras + condiciones + especificaciones."""
    doc = Document()
    for f, h, filas, cond in por_pagina:
        doc.add_paragraph(INSTRUCCION.format(f=f, h=h))
        plano = (['N°', 'Operador', 'Producto'] + ['Peso (g) Muestra evaluada'] * 10 + ['Apariencia'] * 2
                 + ['Observaciones', 'Realizado por:'])
        _tabla(doc, [plano, ENC_TQ[1]] + filas)
        _tabla(doc, cond)
        _tabla(doc, [['Almond Bites - Especificaciones de peso (gramos)', '', '', ''], ['T masa', '99', '', '']])
    doc.sections[0].footer.paragraphs[0].text = 'Revisado por: ______'
    doc.save(ruta)


def test_paginas_separa_cada_bloque_de_la_plantilla(tmp_path):
    ruta = tmp_path / 'p.docx'
    docx_fq32(ruta, [('15/09/2026', '8:00', [fila('Olga', 'Boquita', ['27'])], [['T°C Ambiente', '27', '', '']]),
                     ('____', '____', BLANCAS, COND_TQ_VACIA)])
    pgs = paginas(ruta)
    assert [p['n'] for p in pgs] == [1, 2]
    assert 'Fecha: 15/09/2026' in pgs[0]['instruccion']
    assert pgs[0]['muestras'][2][:4] == ['1', 'Olga', 'Boquita', '27']
    assert pgs[0]['condiciones'] == [['T°C Ambiente', '27', '', '']]   # sin la tabla de especificaciones
    assert pgs[0]['pie'] == 'Revisado por: ______'
    assert len(pgs[1]['muestras']) == 10


def test_extraer_lee_los_tres_documentos(tmp_path):
    archivo_tq, archivo_em, archivo_gf = ARCHIVOS
    docx_fq32(tmp_path / archivo_tq, [('15/09/2026', '8:00', [fila('Olga', 'Boquita', ['27'])], COND_TQ_VACIA),
                                      ('____', '____', BLANCAS, COND_TQ_VACIA)])
    docx_fq32(tmp_path / archivo_gf, [('____', '____', BLANCAS, COND_TQ_VACIA)])
    registros, rechazos = extraer(tmp_path, HOY)
    assert [(r['formato'], r['fecha'], r['fuente']) for r in registros] == [
        ('fq32-tequenos', '2026-09-15', f'{archivo_tq} › página 1')]
    assert rechazos == [{'fuente': archivo_em, 'motivo': 'archivo no encontrado', 'contenido': ''}]
