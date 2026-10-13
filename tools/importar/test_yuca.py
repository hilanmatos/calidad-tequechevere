"""Pruebas del extractor de Yuca Bites y carimañola (extractores/yuca.py) con un libro sintético."""
from datetime import date, datetime, time

from openpyxl import Workbook

from extractores import yuca
from extractores._bites import resolver_producto

HOY = date(2026, 10, 12)


def test_productos_del_excel_al_catalogo():
    r = lambda n, p: resolver_producto(n, p, 'cq-l5-yuca', yuca.REGLAS)  # noqa: E731
    assert r('Yuca bites de queso', '16 UNIDADES ') == 'Yuca Bite de Queso'
    assert r('Yuca bites de Pollo', '16 UNIDADES ') == 'Yuca Bite de Pollo'
    assert r('Yuca bites de carne', '50 UNIDADES') == 'Yuca Bite de Carne'
    assert r('carimañola de carne', '50 UNIDADES') == 'Mini Carimañola'
    assert r('carimañola de carne', '16 UNIDADES') is None   # la carimañola solo va x50
    assert r('Yuky snack de queso blanco', '32 unidades') is None


H1 = {1: datetime(2026, 4, 15), 2: 'Yuca bites de pollo', 3: 63.49, 4: '28 min', 5: 100, 6: '22 pulgadas', 7: 29.4,
      8: 29, 9: 56, 10: '-', 11: 'Correcta', 12: -2.2, 13: time(1, 30), 14: '-', 15: time(1, 20), 16: 27.86,
      17: 5.1, 18: 'Correcta', 19: 4, 20: '-', 21: '-', 22: ' 200°C horno', 23: '6 min', 24: 44, 25: 1, 26: -5,
      27: 'Correcta', 28: '16 UNIDADES ', 29: 'PY012(27103)', 30: '14/10/2026', 32: 0.427, 33: 0.4, 34: 0.027,
      35: 'Ausencia ', 36: 'Correcto', 38: 'Equipo de freido dañado'}
H1_SIN_LOTE = {1: datetime(2025, 12, 11), 2: 'Yuca bites de queso', 3: 56.43, 4: 20, 5: 100, 6: 20,
               28: '32 unidades', 30: datetime(2026, 6, 11), 32: 871, 35: 'Ausencia '}


def _libro(tmp_path, hoja1, anterior=(), tiempos=()):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Hoja1'
    for col, t in yuca.ENCABEZADOS.items():
        ws.cell(row=4, column=col, value=t)
    for i, fila in enumerate(hoja1, start=5):
        for col, v in fila.items():
            ws.cell(row=i, column=col, value=v)
    ant = wb.create_sheet('YUCA BITES')
    ant.cell(row=2, column=2, value='FECHA DE PRODUCCIÓN ')
    ant.cell(row=2, column=4, value='PRESENTACIÓN ')
    ant.cell(row=2, column=5, value='LOTE')
    for i, fila in enumerate(anterior, start=3):
        for col, v in fila.items():
            ant.cell(row=i, column=col, value=v)
    ant.cell(row=3 + len(anterior), column=1, value=99)  # fila de plantilla: solo el N°
    h2 = wb.create_sheet('Hoja2')
    h2.cell(row=1, column=2, value='Tiempos de espera Yuca Bites')
    for i, fila in enumerate(tiempos, start=4):
        for col, v in fila.items():
            h2.cell(row=i, column=col, value=v)
    wb.save(tmp_path / yuca.ARCHIVOS[0])
    return tmp_path


def test_extraer_yuca_mapea_la_fila(tmp_path):
    registros, rechazos = yuca.extraer(_libro(tmp_path, [H1]), HOY)
    assert rechazos == []
    r = registros[0]
    assert (r['formato'], r['fecha'], r['producto'], r['lote']) == ('cq-l5-yuca', '2026-04-15',
                                                                     'Yuca Bite de Pollo', '27103')
    assert r['valores'] == {
        'hYuca': 63.49, 'tiempoCoccion': 28.0, 'tCoccion': 100.0, 'tMasa': 29.4, 'tAmbiente': 29.0,
        'hrAmbiente': 56.0, 'aparienciaMasa': 'Correcta', 'tRelleno': -2.2, 'tiempoEspera': 90.0,
        'tiempoFormado': 80.0, 'pesoPromedio': 27.86, 'pesoRelleno': 5.1, 'aparienciaFormado': 'Correcta',
        'pctNoConformeFormado': 4.0, 'tAceite': 200.0, 'tiempoFritura': 360.0, 'tInternaFritura': 44.0,
        'pctNoConformeFritura': 1.0, 'tInternaCong': -5.0, 'aparienciaCong': 'Correcta',
        'presentacion': '16 UNIDADES', 'vencimiento': '2026-10-14', 'pesoEmpaque': 0.427, 'metales': 'Ausencia',
        'sellado': 'Correcto',
    }
    assert r['observaciones'] == 'Cantidad de agua: 22 pulgadas; Sabor (sensorial): Equipo de freido dañado'


def test_extraer_yuca_hoja_anterior_duplicada_se_rechaza(tmp_path):
    anterior = [
        {1: 1, 2: datetime(2026, 4, 15), 3: 'Yuca bite pollo', 4: '16 unidades', 5: '[PY012] 27103'},
        {1: 2, 2: datetime(2025, 11, 12), 3: 'Yuky snack de queso blanco', 4: '32 unidades'},  # sin lote
        {1: 3, 2: '20/12/2025', 3: 'Yuca bite queso blanco', 4: '16 unidades', 5: '[PY001] 25999'},
    ]
    registros, rechazos = yuca.extraer(_libro(tmp_path, [H1, H1_SIN_LOTE], anterior), HOY)
    assert len(registros) == 2
    por_fila = {x['fuente'].split(' › ', 1)[1]: x for x in rechazos}
    assert por_fila['YUCA BITES › fila 3']['motivo'] == 'hoja anterior duplicada (la producción ya está en Hoja1)'
    assert por_fila['YUCA BITES › fila 3']['contenido'].startswith('misma producción en Hoja1 › fila 5 | ')
    assert por_fila['YUCA BITES › fila 4']['contenido'].startswith('misma producción en Hoja1 › fila 6 | ')
    assert por_fila['YUCA BITES › fila 5']['motivo'].startswith('hoja anterior sin equivalente')
    assert len(rechazos) == 3  # la fila de plantilla con solo el N° no cuenta


def test_extraer_yuca_hoja2_estudio_de_tiempos_va_a_rechazos(tmp_path):
    tiempos = [{1: 'carrito 1', 2: time(9, 35), 4: -26.37}]
    _, rechazos = yuca.extraer(_libro(tmp_path, [H1], tiempos=tiempos), HOY)
    assert len(rechazos) == 1
    assert rechazos[0]['motivo'].startswith('estudio de tiempos')
    assert rechazos[0]['fuente'].endswith('Hoja2 › fila 4')
