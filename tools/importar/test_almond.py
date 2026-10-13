"""Pruebas del extractor de Almond Bites (extractores/almond.py) con un libro sintético."""
from datetime import date, datetime, time

from openpyxl import Workbook

from extractores import almond
from extractores._bites import resolver_producto

HOY = date(2026, 10, 12)


def test_productos_del_excel_al_catalogo():
    r = lambda n, p='20 UNIDADES': resolver_producto(n, p, 'cq-l4-almond', almond.REGLAS)  # noqa: E731
    assert r('Almond bites con queso') == 'Almond Bite de Queso'
    assert r('Almond bites con chocolate') == 'Almond Bite de Chocolate'
    assert r('Almond bites con chocolate', '-') == 'Almond Bite de Chocolate'
    assert r('Almond bites con guayaba') is None


def _libro(tmp_path, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Control de proceso'
    ws.cell(row=1, column=1, value='LINEA 4-ALMOND BITES')
    for col, t in almond.ENCABEZADOS.items():
        ws.cell(row=3, column=col, value=t)
    for i, fila in enumerate(filas, start=4):
        for col, v in fila.items():
            ws.cell(row=i, column=col, value=v)
    wb.save(tmp_path / almond.ARCHIVOS[0])
    return tmp_path


BASE = {1: '16/04/2026', 2: 'Almond bites con queso', 3: 33.2, 4: 28, 5: 56, 6: '-', 7: 'Correcta', 8: 9.7,
        9: '1 hora', 10: 14.8, 11: time(0, 45), 12: 15.9, 13: 4.95, 14: 'Correcta', 15: 2, 16: '-', 17: '-',
        18: '20 UNIDADES', 19: 'PK006(27114)', 20: '16/10/26', 21: 0.403, 22: 0.32, 23: 0.083, 25: 'Correcto',
        26: 'Correcto'}


def test_extraer_almond_mapea_la_fila(tmp_path):
    registros, rechazos = almond.extraer(_libro(tmp_path, [BASE]), HOY)
    assert rechazos == []
    r = registros[0]
    assert (r['formato'], r['fecha'], r['producto'], r['lote']) == ('cq-l4-almond', '2026-04-16',
                                                                     'Almond Bite de Queso', '27114')
    assert r['valores'] == {
        'tMasa': 33.2, 'tAmbiente': 28.0, 'hrAmbiente': 56.0, 'aparienciaMasa': 'Correcta', 'tRelleno': 9.7,
        'tiempoEspera': 60.0, 'tRellenoFinal': 14.8, 'tiempoFormado': 45.0, 'pesoPromedio': 15.9,
        'pesoRelleno': 4.95, 'aparienciaFormado': 'Correcta', 'pctNoConforme': 2.0, 'presentacion': '20 UNIDADES',
        'vencimiento': '2026-10-16', 'pesoEmpaque': 0.403, 'sellado': 'Correcto', 'impresionLote': 'Correcta',
    }
    assert r['fuente'] == 'CONTROL DE CALIDAD ALMOND BITES.xlsx › Control de proceso › fila 4'


def test_extraer_almond_textos_y_notas_sin_encabezado(tmp_path):
    fila = {**BASE, 14: 'La mayoría abiertas', 28: 'Correcto', 29: 'grisáceo',
            34: 'Se realizó a las 6:30 am, no se le pudo dar seguimiento '}
    registros, _ = almond.extraer(_libro(tmp_path, [fila]), HOY)
    r = registros[0]
    # lo descriptivo va a observaciones; "Correcto" en la escala es Buena (Operaciones, 12/10/2026)
    assert 'aparienciaFormado' not in r['valores'] and r['valores']['sabor'] == 'Buena'
    assert r['observaciones'] == ('Apariencia (formado): La mayoría abiertas; '
                                  'Color (sensorial): grisáceo; '
                                  'Se realizó a las 6:30 am, no se le pudo dar seguimiento')


def test_extraer_almond_aplica_unicidad_por_lote(tmp_path):
    otra = {**BASE, 1: datetime(2026, 5, 5), 19: 'PK006 (27114)', 3: 35.4}
    registros, rechazos = almond.extraer(_libro(tmp_path, [BASE, BASE, otra]), HOY)
    assert len(registros) == 1
    assert [x['motivo'].split(' (')[0] for x in rechazos] == ['duplicado exacto', 'lote repetido con otros datos']
